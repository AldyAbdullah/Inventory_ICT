import pandas as pd
import io
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_report(data_barang, data_transaksi):
    output = io.BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        waktu_cetak = (datetime.utcnow() + timedelta(hours=8)).strftime('%Y-%m-%d %H:%M:%S WITA')
        
        # Konfigurasi Gaya (Styling)
        font_title = Font(name='Segoe UI', size=14, bold=True, color='0A2540')
        font_subtitle = Font(name='Segoe UI', size=11, bold=True, color='334155')
        font_meta = Font(name='Segoe UI', size=9, italic=True, color='64748B')
        font_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
        fill_header = PatternFill(start_color='0A2540', end_color='0A2540', fill_type='solid')
        font_data = Font(name='Segoe UI', size=10, color='1E293B')
        
        border_thin = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        
        align_center = Alignment(horizontal='center', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')

        # --- SHEET 1: MASTER STOK IT ---
        list_barang = []
        for idx, item in enumerate(data_barang, start=1):
            status_stok = 'Habis' if item.stok == 0 else ('Kritis' if item.stok <= 5 else 'Aman')
            list_barang.append({
                'No': idx,
                'Kode Barang': item.kode_barang,
                'Nama Barang': item.nama_barang,
                'Kategori': item.kategori,
                'Stok': item.stok,
                'Satuan': item.satuan,
                'Status': status_stok
            })
            
        df_barang = pd.DataFrame(list_barang)
        df_barang.to_excel(writer, index=False, sheet_name='Master Stok IT', startrow=4)
        
        ws_barang = writer.sheets['Master Stok IT']
        ws_barang['A1'] = "JOB Pertamina-Medco E&P Tomori"
        ws_barang['A1'].font = font_title
        ws_barang['A2'] = "Laporan Master Stok Barang IT Habis Pakai"
        ws_barang['A2'].font = font_subtitle
        ws_barang['A3'] = f"Tanggal Cetak: {waktu_cetak}"
        ws_barang['A3'].font = font_meta

        for col in range(1, len(df_barang.columns) + 1):
            cell = ws_barang.cell(row=5, column=col)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_thin

        for row in range(6, len(df_barang) + 6):
            for col in range(1, len(df_barang.columns) + 1):
                cell = ws_barang.cell(row=row, column=col)
                cell.font = font_data
                cell.border = border_thin
                if col in [1, 2, 5, 6, 7]:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        # --- SHEET 2: RIWAYAT TRANSAKSI ---
        list_transaksi = []
        for idx, log in enumerate(data_transaksi, start=1):
            list_transaksi.append({
                'No': idx,
                'Waktu (UTC)': log.tanggal.strftime('%Y-%m-%d %H:%M:%S'),
                'Kode Barang': log.barang.kode_barang,
                'Nama Barang': log.barang.nama_barang,
                'Jenis': log.jenis,
                'Jumlah': log.jumlah,
                'Satuan': log.barang.satuan,
                'Keterangan / Referensi': log.keterangan
            })
            
        df_transaksi = pd.DataFrame(list_transaksi)
        df_transaksi.to_excel(writer, index=False, sheet_name='Riwayat Transaksi', startrow=4)
        
        ws_transaksi = writer.sheets['Riwayat Transaksi']
        ws_transaksi['A1'] = "JOB Pertamina-Medco E&P Tomori"
        ws_transaksi['A1'].font = font_title
        ws_transaksi['A2'] = "Laporan Riwayat Mutasi Barang IT"
        ws_transaksi['A2'].font = font_subtitle
        ws_transaksi['A3'] = f"Tanggal Cetak: {waktu_cetak}"
        ws_transaksi['A3'].font = font_meta

        for col in range(1, len(df_transaksi.columns) + 1):
            cell = ws_transaksi.cell(row=5, column=col)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_thin

        for row in range(6, len(df_transaksi) + 6):
            for col in range(1, len(df_transaksi.columns) + 1):
                cell = ws_transaksi.cell(row=row, column=col)
                cell.font = font_data
                cell.border = border_thin
                if col in [1, 2, 3, 5, 6, 7]:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        for ws in [ws_barang, ws_transaksi]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row >= 5 and cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output.seek(0)
    tanggal_hari_ini = (datetime.utcnow() + timedelta(hours=8)).strftime('%Y-%m-%d')
    nama_file = f'Laporan_Inventaris_JOB_Tomori_{tanggal_hari_ini}.xlsx'
    
    return output, nama_file