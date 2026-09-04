from flask import Flask, render_template, request, redirect, url_for
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
from flask import abort

app = Flask(__name__)
app.secret_key = 'kunci_rahasia_inventaris_job_tomori_sangat_aman'

app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///inventaris_job_tomori.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
db = SQLAlchemy(app)

# Konfigurasi Flask-Login
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login' # Halaman tujuan jika user belum login

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if current_user.role != 'Admin':
            abort(403) # Mengembalikan error 403 Forbidden jika bukan Admin
        return f(*args, **kwargs)
    return decorated_function

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

# Model Tabel User
class User(UserMixin, db.Model):
    __tablename__ = 'user'
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(50), unique=True, nullable=False)
    password = db.Column(db.String(200), nullable=False)
    role = db.Column(db.String(20), default='Admin') # 'Admin' atau 'Viewer'

class BarangIT(db.Model):
    __tablename__ = 'barang_it' # Deklarasi nama tabel secara eksplisit
    id = db.Column(db.Integer, primary_key=True)
    kode_barang = db.Column(db.String(20), unique=True, nullable=False)
    nama_barang = db.Column(db.String(100), nullable=False)
    kategori = db.Column(db.String(50), nullable=False)
    stok = db.Column(db.Integer, default=0)
    satuan = db.Column(db.String(20), nullable=False)

# Tambahkan Model Transaksi
class Transaksi(db.Model):
    __tablename__ = 'transaksi'
    id = db.Column(db.Integer, primary_key=True)
    barang_id = db.Column(db.Integer, db.ForeignKey('barang_it.id'), nullable=False)
    jenis = db.Column(db.String(10), nullable=False) # 'Masuk' atau 'Keluar'
    jumlah = db.Column(db.Integer, nullable=False)
    tanggal = db.Column(db.DateTime, default=datetime.utcnow)
    keterangan = db.Column(db.String(200))

    # Relasi untuk memanggil data barang dari tabel transaksi
    barang = db.relationship('BarangIT', backref=db.backref('riwayat_transaksi', lazy=True))

with app.app_context():
    db.create_all()

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        
        user = User.query.filter_by(username=username).first()
        if user and check_password_hash(user.password, password):
            login_user(user)
            return redirect(url_for('index'))
        else:
            return render_template('login.html', error="Username atau password salah!")
            
    return render_template('login.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))

@app.route('/')
@login_required
def index():
    # Menangkap parameter pencarian dari URL
    search_query = request.args.get('search', '')
    
    # Query data untuk tabel (berdasarkan pencarian jika ada)
    if search_query:
        barang_habis_pakai = BarangIT.query.filter(
            (BarangIT.nama_barang.ilike(f'%{search_query}%')) | 
            (BarangIT.kode_barang.ilike(f'%{search_query}%'))
        ).all()
    else:
        barang_habis_pakai = BarangIT.query.all()
        
    # Kalkulasi data statistik (menggunakan seluruh data di database, bukan data hasil filter)
    semua_barang = BarangIT.query.all()
    total_jenis_barang = len(semua_barang)
    total_stok = sum(item.stok for item in semua_barang)
    stok_kritis = sum(1 for item in semua_barang if item.stok <= 5)

    return render_template('index.html', 
                           data=barang_habis_pakai,
                           total_jenis=total_jenis_barang,
                           total_stok=total_stok,
                           stok_kritis=stok_kritis,
                           search_query=search_query)

@app.route('/tambah', methods=['GET', 'POST'])
@login_required
@admin_required
def tambah():
    # Jika user menekan tombol 'Simpan' (metode POST)
    if request.method == 'POST':
        # Mengambil data dari form HTML
        kode = request.form['kode_barang']
        nama = request.form['nama_barang']
        kategori = request.form['kategori']
        stok = request.form['stok']
        satuan = request.form['satuan']

        # Membuat objek barang baru
        barang_baru = BarangIT(
            kode_barang=kode, 
            nama_barang=nama, 
            kategori=kategori, 
            stok=stok, 
            satuan=satuan
        )
        
        # Menyimpan ke database
        db.session.add(barang_baru)
        db.session.commit()
        
        # Mengarahkan kembali ke halaman utama setelah berhasil
        return redirect(url_for('index'))
    
    # Jika user hanya membuka halaman (metode GET), tampilkan form
    return render_template('tambah.html')

@app.route('/edit/<int:id>', methods=['GET', 'POST'])
@login_required
@admin_required
def edit(id):
    # Mencari data barang berdasarkan ID, jika tidak ada kembalikan error 404
    barang = BarangIT.query.get_or_404(id)
    
    if request.method == 'POST':
        # Memperbarui data objek barang dengan data dari form
        barang.kode_barang = request.form['kode_barang']
        barang.nama_barang = request.form['nama_barang']
        barang.kategori = request.form['kategori']
        barang.stok = request.form['stok']
        barang.satuan = request.form['satuan']
        
        # Menyimpan perubahan ke database
        db.session.commit()
        return redirect(url_for('index'))
        
    return render_template('edit.html', barang=barang)

@app.route('/hapus/<int:id>')
@login_required
@admin_required
def hapus(id):
    barang = BarangIT.query.get_or_404(id)
    db.session.delete(barang)
    db.session.commit()
    return redirect(url_for('index'))

@app.route('/transaksi/<int:id>', methods=['GET', 'POST'])
@login_required
@admin_required
def catat_transaksi(id):
    barang = BarangIT.query.get_or_404(id)
    
    if request.method == 'POST':
        jenis = request.form['jenis']
        jumlah = int(request.form['jumlah'])
        keterangan = request.form['keterangan']

        # Validasi agar stok tidak minus
        if jenis == 'Keluar' and jumlah > barang.stok:
            return "Error: Stok tidak mencukupi untuk dikeluarkan.", 400

        # Catat ke tabel Transaksi
        transaksi_baru = Transaksi(
            barang_id=barang.id, 
            jenis=jenis, 
            jumlah=jumlah, 
            keterangan=keterangan
        )
        
        # Perbarui stok di tabel BarangIT
        if jenis == 'Masuk':
            barang.stok += jumlah
        elif jenis == 'Keluar':
            barang.stok -= jumlah

        db.session.add(transaksi_baru)
        db.session.commit()
        
        return redirect(url_for('index'))
        
    return render_template('form_transaksi.html', barang=barang)

@app.route('/riwayat')
@login_required
def riwayat():
    # Mengambil data transaksi diurutkan dari yang terbaru
    data_riwayat = Transaksi.query.order_by(Transaksi.tanggal.desc()).all()
    return render_template('riwayat.html', riwayat=data_riwayat)

import pandas as pd
from flask import send_file
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

@app.route('/export/excel')
@login_required
def export_excel():
    output = io.BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        waktu_cetak = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        
        # Konfigurasi Gaya (Styling) openpyxl
        font_title = Font(name='Segoe UI', size=14, bold=True, color='0A2540')
        font_subtitle = Font(name='Segoe UI', size=11, bold=True, color='334155')
        font_meta = Font(name='Segoe UI', size=9, italic=True, color='64748B')
        
        font_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
        fill_header = PatternFill(start_color='0A2540', end_color='0A2540', fill_type='solid')
        
        font_data = Font(name='Segoe UI', size=10, color='1E293B')
        border_thin = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        
        align_center = Alignment(horizontal='center', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')
        align_right = Alignment(horizontal='right', vertical='center')

        # --- SHEET 1: MASTER STOK IT ---
        data_barang = BarangIT.query.all()
        list_barang = []
        for idx, item in enumerate(data_barang, start=1):
            status_stok = 'Habis' if item.stok == 0 else ('Kritis' if item.stok <= 5 else 'Aman')
            list_barang.append({
                'No': idx,
                'Kode Barang': item.kode_barang,
                'Nama Barang': item.nama_barang,
                'Kategori': item.kategori,
                'Stok': item.stok,
                'Satuan': item.satuan,
                'Status': status_stok
            })
        df_barang = pd.DataFrame(list_barang)
        df_barang.to_excel(writer, index=False, sheet_name='Master Stok IT', startrow=4)
        
        ws_barang = writer.sheets['Master Stok IT']
        ws_barang['A1'] = "JOB Pertamina-Medco E&P Tomori"
        ws_barang['A1'].font = font_title
        ws_barang['A2'] = "Laporan Master Stok Barang IT Habis Pakai"
        ws_barang['A2'].font = font_subtitle
        ws_barang['A3'] = f"Tanggal Cetak: {waktu_cetak}"
        ws_barang['A3'].font = font_meta

        # Styling Header Tabel Master
        for col in range(1, len(df_barang.columns) + 1):
            cell = ws_barang.cell(row=5, column=col)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_thin

        # Styling Data Master
        for row in range(6, len(df_barang) + 6):
            for col in range(1, len(df_barang.columns) + 1):
                cell = ws_barang.cell(row=row, column=col)
                cell.font = font_data
                cell.border = border_thin
                # Perataan teks berdasarkan kolom
                if col in [1, 2, 5, 6, 7]:  # No, Kode, Stok, Satuan, Status
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        # --- SHEET 2: RIWAYAT TRANSAKSI ---
        data_transaksi = Transaksi.query.order_by(Transaksi.tanggal.desc()).all()
        list_transaksi = []
        for idx, log in enumerate(data_transaksi, start=1):
            list_transaksi.append({
                'No': idx,
                'Waktu (UTC)': log.tanggal.strftime('%Y-%m-%d %H:%M:%S'),
                'Kode Barang': log.barang.kode_barang,
                'Nama Barang': log.barang.nama_barang,
                'Jenis': log.jenis,
                'Jumlah': log.jumlah,
                'Satuan': log.barang.satuan,
                'Keterangan / Referensi': log.keterangan
            })
        df_transaksi = pd.DataFrame(list_transaksi)
        df_transaksi.to_excel(writer, index=False, sheet_name='Riwayat Transaksi', startrow=4)
        
        ws_transaksi = writer.sheets['Riwayat Transaksi']
        ws_transaksi['A1'] = "JOB Pertamina-Medco E&P Tomori"
        ws_transaksi['A1'].font = font_title
        ws_transaksi['A2'] = "Laporan Riwayat Mutasi Barang IT"
        ws_transaksi['A2'].font = font_subtitle
        ws_transaksi['A3'] = f"Tanggal Cetak: {waktu_cetak}"
        ws_transaksi['A3'].font = font_meta

        # Styling Header Tabel Transaksi
        for col in range(1, len(df_transaksi.columns) + 1):
            cell = ws_transaksi.cell(row=5, column=col)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_thin

        # Styling Data Transaksi
        for row in range(6, len(df_transaksi) + 6):
            for col in range(1, len(df_transaksi.columns) + 1):
                cell = ws_transaksi.cell(row=row, column=col)
                cell.font = font_data
                cell.border = border_thin
                if col in [1, 2, 3, 5, 6, 7]:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        # Auto-adjust Column Width untuk kedua Sheet
        for ws in [ws_barang, ws_transaksi]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    # Hitung panjang teks hanya dari baris data ke bawah agar tidak terpengaruh judul yang panjang
                    if cell.row >= 5 and cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output.seek(0)
    tanggal_hari_ini = datetime.now().strftime('%Y-%m-%d')
    nama_file = f'Laporan_Inventaris_JOB_Tomori_{tanggal_hari_ini}.xlsx'
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=nama_file
    )
@app.route('/users')
@login_required
@admin_required
def kelola_user():
    users = User.query.all()
    return render_template('users.html', users=users)

@app.route('/users/tambah', methods=['POST'])
@login_required
@admin_required
def tambah_user():
    username = request.form['username']
    password = request.form['password']
    role = request.form['role']

    # Validasi jika username sudah digunakan
    user_exists = User.query.filter_by(username=username).first()
    if user_exists:
        return "Error: Username sudah digunakan.", 400

    hashed_password = generate_password_hash(password)
    user_baru = User(username=username, password=hashed_password, role=role)
    db.session.add(user_baru)
    db.session.commit()
    
    return redirect(url_for('kelola_user'))

@app.route('/users/hapus/<int:id>')
@login_required
@admin_required
def hapus_user(id):
    # Mencegah Admin menghapus akun yang sedang digunakannya sendiri
    if id == current_user.id:
        return "Error: Tidak dapat menghapus akun yang sedang digunakan.", 400
        
    user = User.query.get_or_404(id)
    db.session.delete(user)
    db.session.commit()
    
    return redirect(url_for('kelola_user'))

if __name__ == '__main__':
    app.run(debug=True)